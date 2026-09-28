"""Local static/time-series results, independent of analysis software and anatomy."""
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4
import hashlib
import re
import unicodedata
import numpy as np
from .result_display import COLORMAPS,metric_display,marker_options,spatial_options,visible_values
from .contact_matching import ContactNameIndex


@dataclass
class AnalysisResult:
    name: str
    kind: str
    channels: list[str]
    metrics: list[str]
    values: np.ndarray  # channel, frame, metric; NaN means missing, never zero
    times: np.ndarray
    units: list[str]
    time_unit: str = ''
    uid: str = field(default_factory=lambda: 'result_' + uuid4().hex)
    provenance: dict = field(default_factory=dict)
    # Original channel -> one contact UID, or two UIDs for a bipolar midpoint.
    bindings: dict = field(default_factory=dict)
    binding_modes: dict = field(default_factory=dict)
    aliases: dict = field(default_factory=dict)
    settings: dict = field(default_factory=dict)


def fail(text):
    from .imaging import InputError
    raise InputError(text)


def normalized(name):
    return unicodedata.normalize('NFKC', str(name)).strip().replace('−','-').replace('–','-').casefold()


def validate_result(result):
    n, t, m = len(result.channels), len(result.times), len(result.metrics)
    if (result.kind not in ('static','time_series') or not n or not t or not m
            or result.values.shape != (n,t,m) or len(result.units) != m
            or result.times.shape != (t,) or np.isinf(result.values).any()
            or not np.isfinite(result.times).all()
            or (result.kind == 'static' and t != 1)
            or not result.uid.startswith('result_') or not result.uid[7:].isalnum()
            or not result.name.strip() or any(not str(x).strip() for x in result.channels + result.metrics)
            or len({normalized(x) for x in result.channels}) != n
            or len(set(result.metrics)) != m or result.time_unit not in ('','ms','s')):
        fail('解析結果のチャンネル・値・時間情報が不正です。')
    if any(k not in result.channels or not isinstance(v,list) or len(v) not in (1,2)
           or any(not isinstance(x,str) or not x for x in v) or len(set(v)) != len(v)
           for k,v in result.bindings.items()):
        fail('解析チャンネルとコンタクトの対応情報が不正です。')
    if any(k not in result.channels or not isinstance(v,str) for k,v in result.aliases.items()):
        fail('解析チャンネルの対応名が不正です。')
    if any(k not in result.channels or v not in ('monopolar','bipolar') for k,v in result.binding_modes.items()):
        fail('解析チャンネルとコンタクトの対応情報が不正です。')


def number(value, coordinate):
    if value is None or (isinstance(value,str) and value.strip().lower() in ('','na','n/a','nan')):
        return np.nan
    if isinstance(value,bool): fail(f'数値列に真偽値があります：{coordinate}')
    try: value = float(value)
    except (TypeError, ValueError): fail(f'数値として読めないセルがあります：{coordinate}')
    if not np.isfinite(value): fail(f'有限の数値ではないセルがあります：{coordinate}')
    return value


def read_workbook(path,time_unit=''):
    """pyHFO Channel Ranking, or a generic channel/metric or channel/time matrix.

    Read only cached Excel values, no macros, formulas or external links executed.
    A separate ranking table after a blank header is not mistaken for new metrics.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    path = Path(path)
    if time_unit not in ('','ms','s'):fail('時間単位を確認してください。')
    if path.suffix.lower() != '.xlsx': fail('解析結果は .xlsx 形式で指定してください。')
    try:
        workbook = load_workbook(path,read_only=True,data_only=False,keep_links=False)
    except Exception: fail('Excelを読み込めませんでした。.xlsx形式を確認してください。')
    cached = None
    try:
        hfo = 'Channel Ranking' in workbook.sheetnames
        sheet = workbook['Channel Ranking'] if hfo else workbook.worksheets[0]
        if sheet.max_row > 10001 or sheet.max_column > 20001 or sheet.max_row*sheet.max_column > 5_000_000:
            fail('解析表が大きすぎます。チャンネル数または時間列を分けてください。')
        cells = list(sheet.iter_rows())
        if not cells or len(cells) < 2: fail('解析表にデータ行がありません。')
        has_formula = any(c.data_type == 'f' for row in cells for c in row)
        if has_formula:
            cached = load_workbook(path,read_only=True,data_only=True,keep_links=False)
            cache_rows = list(cached[sheet.title].values)
        rows = []
        for i,row in enumerate(cells):
            values=[]
            for j,cell in enumerate(row):
                value = cache_rows[i][j] if cell.data_type == 'f' else cell.value
                if cell.data_type == 'e' or (cell.data_type == 'f' and value is None):
                    fail(f'Excelの計算結果を確認してください：{cell.coordinate}')
                values.append(value)
            rows.append(values)
        header = rows[0]
        # A labelled channel column may follow an index column. Unlabelled
        # metadata before the metric block is ignored; a second ranking table
        # after its blank separator is never joined by row order.
        channel_col=next((i for i,x in enumerate(header) if normalized(str(x)) in ('channel','channel_name','channel name')),0)
        start=next((i for i in range(channel_col+1,len(header)) if header[i] is not None and str(header[i]).strip()),len(header))
        end = next((i for i in range(start,len(header)) if header[i] is None or str(header[i]).strip()==''),len(header))
        if end <= start: fail('1行目に指標名または時間、1列目にチャンネル名が必要です。')
        numeric_headers=[]
        for x in header[start:end]:
            try: numeric_headers.append(float(x))
            except (TypeError, ValueError): break
        dynamic = len(numeric_headers)==end-start and not hfo
        channels=[]; data=[]; source_rows=[]
        for i,row in enumerate(rows[1:],2):
            if row[channel_col] is None or not str(row[channel_col]).strip():
                if any(v is not None for v in row[start:end]): fail(f'チャンネル名がないデータ行があります：{i}')
                continue
            channels.append(str(row[channel_col]).strip()); source_rows.append(i)
            data.append([number(row[j],f'{get_column_letter(j+1)}{i}') for j in range(start,end)])
        array=np.asarray(data,dtype=np.float64)
        if not channels: fail('解析表にチャンネルがありません。')
        if dynamic:
            metrics=['gamma' if any(x in path.stem.lower() for x in ('audio','photo','gamma')) else 'value']
            times=np.asarray(numeric_headers,dtype=float); values=array[:,:,None]; units=['']
            # Never infer ms or percentage from the numerical range alone.
            condition=re.search(r'(audio|photo)_(stim|res)$',path.stem)
            name=('Gamma '+condition[0] if condition else 'Gamma' if path.stem.lower()=='gamma' else 'Gamma '+path.stem) if metrics==['gamma'] else path.stem
        else:
            metrics=[str(x).strip() for x in header[start:end]]; times=np.array([0.]); values=array[:,None,:]
            units=['events' if hfo else '' for _ in metrics]; name='HFO' if hfo else path.stem
            if hfo and (np.any(values[np.isfinite(values)] < 0) or np.any(values[np.isfinite(values)] % 1)):
                fail('HFOの集計列に負数または整数以外の値があります。')
        result=AnalysisResult(name,'time_series' if dynamic else 'static',channels,metrics,values,times,units,
            time_unit=time_unit if dynamic else '',
            provenance={'adapter':'pyhfo-ranking-v1' if hfo else 'channel-matrix-v1',
                'file_name':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                'sheet':sheet.title,'rows':source_rows,'columns':list(range(start+1,end+1)),
                'channel_column':channel_col+1,'time_unit_source':'specified_on_import' if dynamic and time_unit else 'unspecified',
                'original_headers':[str(x) for x in header[start:end]],
                'imported_at':datetime.now(timezone.utc).isoformat(), 'missing_values':int(np.isnan(values).sum())})
        validate_result(result)
        return result
    finally:
        workbook.close()
        if cached: cached.close()


def frame_range(result):
    last=len(result.times)-1
    start=int(np.clip(result.settings.get('start',0),0,last))
    end=int(np.clip(result.settings.get('end',last),start,last))
    return start,end


def time_issue(result):
    if result.kind=='static': return ''
    a,b=frame_range(result); times=result.times[a:b+1]
    if len(times)<2: return '動画には2列以上の時間範囲を選んでください。'
    if np.any(np.diff(times)<=0): return '時間列に重複・逆行があります。元Excelを確認するか、出力範囲を指定してください。'
    return ''


def color_limits(result, metric):
    saved=result.settings.get('limits',{}).get(result.metrics[metric])
    if saved and len(saved)==2 and np.isfinite(saved).all() and saved[1]>saved[0]: return tuple(saved)
    v=result.values[:,:,metric]; finite=v[np.isfinite(v)]
    if not finite.size: return (0.,1.)
    lo,hi=float(finite.min()),float(finite.max())
    if lo<0: return (-max(abs(lo),abs(hi),1e-9),max(abs(lo),abs(hi),1e-9))
    return (0.,max(hi,1.))


def colors(values, limits,palette='auto',reverse=False):
    """Fixed limits for all frames; missing values use gray only in the data view."""
    from matplotlib import colormaps
    lo,hi=limits; finite=np.isfinite(values)
    unit=np.clip((np.where(finite,values,lo)-lo)/(hi-lo),0,1)
    if palette not in COLORMAPS or palette=='auto':palette='coolwarm' if lo<0 else 'inferno'
    rgba=colormaps[palette+('_r' if reverse else '')](unit)
    rgba[~finite]=(.36,.40,.46,1.)
    return rgba[:,:3]


def alias_candidates(result, others):
    """Suggestions only. Never use prefix or row order for automatic spatial binding."""
    # Repeated truncated labels in other time-series files are not independent
    # evidence of a complete channel name. Use static references or mapped rows.
    names={r.aliases.get(c,c) for r in others if r.uid!=result.uid for c in r.channels
           if r.kind=='static' or c in r.bindings}
    complete={normalized(n) for n in names}
    return {c:matches[0] for c in result.channels
            if normalized(c) not in complete
            and len(matches:=[n for n in sorted(names) if normalized(n).startswith(normalized(c))])==1
            and matches[0]!=c}


def exact_bindings(result, contacts):
    names={}
    for c in contacts: names.setdefault(normalized(c.name),[]).append(c.uid)
    output={}
    for original in result.channels:
        name=normalized(result.aliases.get(original,original))
        if len(names.get(name,[]))==1:
            output[original]=names[name].copy(); continue
        candidates=[]
        for i,char in enumerate(name):
            if char!='-': continue
            first,second=names.get(name[:i],[]),names.get(name[i+1:],[])
            if len(first)==len(second)==1 and first!=second: candidates.append(first+second)
        if len(candidates)==1: output[original]=candidates[0]
    return output


def shared_binding_candidates(result,others,aliases):
    """Reuse patient-local bindings only as dialog candidates."""
    output={}
    for original in result.channels:
        name=normalized(aliases.get(original,original)); candidates=set()
        for other in others:
            if other.uid==result.uid: continue
            for channel,ids in other.bindings.items():
                if normalized(other.aliases.get(channel,channel))==name: candidates.add(tuple(ids))
        if len(candidates)==1: output[original]=list(next(iter(candidates)))
    return output


def resolved_positions(result, contacts):
    by_id={c.uid:c for c in contacts}; indices=[]; positions=[]
    name_index=ContactNameIndex(contacts)
    for index,channel in enumerate(result.channels):
        ids=result.bindings.get(channel,[])
        if not ids: continue
        mode=result.binding_modes.get(channel)
        arity=(2 if mode=='bipolar' else 1) if mode else binding_arity(result.aliases.get(channel,channel),contacts,name_index)
        if (len(ids)!=arity
                or len(set(ids))!=len(ids) or any(uid not in by_id for uid in ids)): continue
        position=np.mean([by_id[uid].position for uid in ids],axis=0)
        if not np.isfinite(position).all(): continue
        indices.append(index); positions.append(position)
    return np.asarray(indices,dtype=int),np.asarray(positions,dtype=float).reshape(-1,3)


def binding_arity(name,contacts,name_index=None):
    return (name_index or ContactNameIndex(contacts)).arity(name)


def channel_tokens(name):
    match=re.fullmatch(r'([A-Za-z][A-Za-z_]*)(\d+)-([A-Za-z][A-Za-z_]*)(\d+)',name)
    if match: return [(match[1],int(match[2])),(match[3],int(match[4]))]
    match=re.fullmatch(r'([A-Za-z][A-Za-z_]*)(\d+)',name)
    return [(match[1],int(match[2]))] if match else []


def lead_binding(name,contacts,lead_map):
    tokens=channel_tokens(name)
    if not tokens: return []
    result=[]
    for source,number in tokens:
        if source not in lead_map: return []
        group,reverse=lead_map[source]; candidates=[c for c in contacts if c.group==group]
        by_number={}
        for contact in candidates:
            match=re.search(r'(\d+)$',contact.name)
            if match: by_number.setdefault(int(match[1]),[]).append(contact.uid)
        if not by_number: return []
        index=max(by_number)+1-number if reverse else number
        if len(by_number.get(index,[]))!=1: return []
        result.append(by_number[index][0])
    return result if len(set(result))==len(result) else []


def overlay_points(result,contacts,metric,frame):
    indices,positions=resolved_positions(result,contacts)
    options=metric_display(result,metric)
    values=result.values[indices,frame,metric]; valid=visible_values(values,options)
    return positions[valid],colors(values[valid],color_limits(result,metric),options['colormap'],options['reverse']),indices[valid]


def result_record(result):
    return {k:getattr(result,k) for k in ('name','kind','channels','metrics','units','time_unit','uid','provenance','bindings','binding_modes','aliases','settings')}


def export_record(result,contacts,metric,frames,fps=None,view='channels'):
    indices,positions=resolved_positions(result,contacts)
    return {'schema':'cortex-result-export/1',**result_record(result),'metric':result.metrics[metric],
        'color_limits':list(color_limits(result,metric)),'view':view,'fps':fps,
        'display_options':metric_display(result,metric),'marker_options':marker_options(result),
        'spatial_display':spatial_options(result),
        'visible_channel_indices':{str(i):indices[visible_values(result.values[indices,i,metric],metric_display(result,metric))].tolist() for i in frames},
        'frames':[{'source_column':result.provenance.get('columns',list(range(2,len(result.times)+2)))[i] if result.kind=='time_series' else None,
                   'index':i,'time':float(result.times[i]) if result.kind=='time_series' else None} for i in frames],
        'playback':'one_source_column_per_frame_no_interpolation',
        'values':[[float(v) if np.isfinite(v) else None for v in row]
                  for row in result.values[:,frames,metric]],
        'missing_value':None,
        'coordinate_system':'scanner_RAS_mm','bipolar_display':'midpoint_of_two_contacts',
        'mapped_channels':len(indices),'unmapped_channels':len(result.channels)-len(indices),
        'positions':[{'channel':result.channels[int(i)],'position':p.tolist(),
                      'contact_uids':result.bindings[result.channels[int(i)]]} for i,p in zip(indices,positions)]}
